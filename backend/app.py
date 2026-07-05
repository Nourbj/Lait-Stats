# -*- coding: utf-8 -*-
from flask import Flask, request, jsonify, send_file, make_response
from flask_cors import CORS
import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
import io
from datetime import datetime
import unicodedata
from copy import copy
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
import re
import csv
import traceback
import os
import math

app = Flask(__name__)
CORS(app)

def normalize_str(s):
    if s is None:
        return ""
    s_str = str(s).strip()
    if s_str.startswith('='):
        return ""
    # Remove accents and lowercase
    s_str = s_str.lower()
    s_str = "".join(
        c for c in unicodedata.normalize('NFD', s_str)
        if unicodedata.category(c) != 'Mn'
    )
    return s_str

def is_csv_bytes(file_bytes):
    if file_bytes.startswith(b'PK\x03\x04'):
        return False
    if file_bytes.startswith(b'\xd0\xcf\x11\xe0'):
        return False
    return True

def parse_csv_to_rows(file_bytes):
    text = None
    for encoding in ['utf-8-sig', 'utf-8', 'iso-8859-1']:
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Impossible de dÃ©coder le fichier CSV.")
    
    first_line = text.split('\n')[0] if text else ""
    delimiter = ','
    if ';' in first_line:
        delimiter = ';'
    elif '\t' in first_line:
        delimiter = '\t'
        
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return list(reader)

def load_workbook_from_bytes(file_bytes, data_only=False):
    if is_csv_bytes(file_bytes):
        rows = parse_csv_to_rows(file_bytes)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "CSV Data"
        for r_idx, row in enumerate(rows, 1):
            for c_idx, val in enumerate(row, 1):
                if val is not None:
                    val_str = str(val).strip()
                    if val_str.isdigit():
                        ws.cell(row=r_idx, column=c_idx, value=int(val_str))
                    else:
                        try:
                            ws.cell(row=r_idx, column=c_idx, value=float(val_str.replace(',', '.')))
                        except ValueError:
                            ws.cell(row=r_idx, column=c_idx, value=val_str)
                else:
                    ws.cell(row=r_idx, column=c_idx, value="")
        return wb
    else:
        # Try to load as an openpyxl workbook (xlsx). If that fails (e.g. old .xls
        # BIFF format), fall back to pandas which can read legacy Excel files
        # and then convert the DataFrame(s) into an openpyxl Workbook.
        try:
            return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=data_only)
        except Exception as e:
            try:
                import pandas as _pd
                # read all sheets
                xls = _pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
            except Exception:
                # re-raise original exception for clarity
                raise e

            wb = openpyxl.Workbook()
            first = True
            for sheet_name, df in xls.items():
                if first:
                    ws = wb.active
                    ws.title = str(sheet_name)[:31]
                    first = False
                else:
                    ws = wb.create_sheet(title=str(sheet_name)[:31])

                # write header
                for c_idx, col in enumerate(df.columns, 1):
                    ws.cell(row=1, column=c_idx, value=str(col))

                # write data rows
                for r_idx, row in enumerate(df.itertuples(index=False, name=None), 2):
                    for c_idx, val in enumerate(row, 1):
                        if isinstance(val, float) and math.isnan(val):
                            val = None
                        ws.cell(row=r_idx, column=c_idx, value=val)

            return wb

def format_single_name(name_str):
    if not name_str:
        return ""
    words = name_str.split()
    if not words:
        return ""
    if len(words) == 1:
        return words[0].title()
    
    first_word = words[0]
    first_is_upper = first_word.isupper() and len(first_word) > 1
    others_all_upper = all(w.isupper() for w in words[1:])
    
    if first_is_upper and not others_all_upper:
        idx = 0
        while idx < len(words) and words[idx].isupper():
            idx += 1
        if idx < len(words):
            prenom_part = " ".join(words[idx:])
            nom_part = " ".join(words[:idx])
            return f"{prenom_part.title()} {nom_part.upper()}"
            
    idx = len(words) - 1
    while idx >= 0 and words[idx].isupper():
        idx -= 1
    
    if idx >= 0 and idx < len(words) - 1:
        prenom_part = " ".join(words[:idx+1])
        nom_part = " ".join(words[idx+1:])
        return f"{prenom_part.title()} {nom_part.upper()}"
        
    return f"{words[0].title()} {' '.join(words[1:]).upper()}"

def format_employee_name(prenom_val, nom_val):
    if prenom_val is not None and nom_val is not None:
        p_str = str(prenom_val).strip()
        n_str = str(nom_val).strip()
        if p_str and n_str:
            return f"{p_str.title()} {n_str.upper()}"
        elif p_str:
            return format_single_name(p_str)
        elif n_str:
            return format_single_name(n_str)
    elif nom_val is not None:
        return format_single_name(str(nom_val).strip())
    elif prenom_val is not None:
        return format_single_name(str(prenom_val).strip())
    return ""

def copy_style(src_cell, dst_cell):
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = copy(src_cell.number_format)

def numeric_value(value):
    if value is None:
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if math.isnan(value):
            return 0
        return value
    if isinstance(value, str):
        value_str = value.strip().replace(',', '.')
        if not value_str:
            return 0
        try:
            return float(value_str)
        except ValueError:
            return 0
    return 0

def safe_int_or_float(val):
    if val is None:
        return 0
    if isinstance(val, (int, float)):
        if math.isnan(val) or math.isinf(val):
            return 0
        return int(val) if val == int(val) else val
    try:
        f_val = float(val)
        if math.isnan(f_val) or math.isinf(f_val):
            return 0
        return int(f_val) if f_val == int(f_val) else f_val
    except (ValueError, TypeError):
        return val

def find_header_columns(ws, header_row):
    columns = {}
    for c in range(1, ws.max_column + 1):
        header_norm = normalize_str(ws.cell(row=header_row, column=c).value)
        if header_norm:
            columns[header_norm] = c
    return columns

def find_column_by_header(ws, header_row, accepted_headers):
    accepted = {normalize_str(h) for h in accepted_headers}
    for c in range(1, ws.max_column + 1):
        header_norm = normalize_str(ws.cell(row=header_row, column=c).value)
        if header_norm in accepted:
            return c
    return None

def get_or_create_quantite_column(ws, header_row, after_col):
    quantite_col = find_column_by_header(ws, header_row, {'quantite', 'quantity', 'qty'})
    if quantite_col:
        return quantite_col

    quantite_col = after_col + 1
    ws.insert_cols(quantite_col)
    header_cell = ws.cell(row=header_row, column=quantite_col, value='Quantité')
    copy_style(ws.cell(row=header_row, column=after_col), header_cell)
    header_cell.font = Font(bold=True, size=24)
    header_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.column_dimensions[get_column_letter(quantite_col)].width = 18
    return quantite_col

def employee_key_for_row(ws, row, identity_cols):
    values = []
    for col in identity_cols:
        value = ws.cell(row=row, column=col).value
        values.append(str(value).strip() if value is not None else '')
    return tuple(values)

def process_attendance_quantity_sheet(ws, header_row, emp_rows, emp_col=None, prenom_col=None, emp_cells=None):
    presence_col = find_column_by_header(ws, header_row, {'presence'})
    if not presence_col:
        return False

    header_cols = find_header_columns(ws, header_row)
    identity_cols = []
    for header in ('emp no.', 'emp no', 'matricule.', 'matricule', 'prenom.', 'prenom', 'nom.', 'nom'):
        col = header_cols.get(normalize_str(header))
        if col and col not in identity_cols:
            identity_cols.append(col)

    if not identity_cols:
        return False

    quantite_col = get_or_create_quantite_column(ws, header_row, presence_col)

    for merged_range in list(ws.merged_cells.ranges):
        if (
            merged_range.min_col == quantite_col
            and merged_range.max_col == quantite_col
            and merged_range.min_row > header_row
        ):
            try:
                ws.unmerge_cells(str(merged_range))
            except KeyError:
                try:
                    ws.merged_cells.ranges.remove(merged_range)
                except KeyError:
                    pass
                for row in range(merged_range.min_row, merged_range.max_row + 1):
                    for col in range(merged_range.min_col, merged_range.max_col + 1):
                        if row == merged_range.min_row and col == merged_range.min_col:
                            continue
                        if isinstance(ws._cells.get((row, col)), MergedCell):
                            del ws._cells[(row, col)]

    current_key = None
    current_rows = []
    groups = []
    for row in emp_rows:
        key = employee_key_for_row(ws, row, identity_cols)
        if not any(key):
            continue
        if current_key is None or key == current_key:
            current_key = key
            current_rows.append(row)
        else:
            groups.append(current_rows)
            current_key = key
            current_rows = [row]
    if current_rows:
        groups.append(current_rows)

    for rows in groups:
        total = sum(numeric_value(ws.cell(row=row, column=presence_col).value) for row in rows)
        first_row = rows[0]
        last_row = rows[-1]

        for row in rows:
            cell = ws.cell(row=row, column=quantite_col)
            cell.value = None
            copy_style(ws.cell(row=row, column=presence_col), cell)

        quantity_cell = ws.cell(row=first_row, column=quantite_col, value=safe_int_or_float(total))
        copy_style(ws.cell(row=first_row, column=presence_col), quantity_cell)
        quantity_cell.font = Font(bold=True, size=22)
        quantity_cell.alignment = Alignment(horizontal='center', vertical='center')

        if emp_cells is not None and emp_col is not None:
            emp_key = get_employee_key(ws, first_row, header_row, emp_col, prenom_col)
            emp_cells[emp_key] = quantity_cell.coordinate

        if last_row > first_row:
            ws.merge_cells(start_row=first_row, start_column=quantite_col, end_row=last_row, end_column=quantite_col)

    return True

def detect_structure(ws):
    EMPLOYEE_KEYWORDS = {
        "employe", "employee", "employes", "employees",
        "nom", "name", "agent", "collaborateur", "personnel",
        "user", "salarie", "travailleur", "staff",
        "prenom", "nom et prenom", "nom & prenom", "nom prenom",
        "collaborateurs", "agents", "salaries"
    }

    SUMMARY_KEYWORDS = {
        "total", "totaux", "somme", "sum", "moyenne", "average", "global"
    }

    EXCLUDED_HEADERS = {
        "departement", "service", "role", "poste", "fonction", "adresse",
        "telephone", "tel", "email", "mail", "status", "embauche", "contrat",
        "commentaire", "comment", "note", "observation", "remarque", "sexe",
        "genre", "age", "date", "jour", "month", "mois", "annee", "year",
        "temps", "time", "timestamp", "horodatage", "prix", "price", "valeur", "montant"
    }

    max_r = min(30, ws.max_row)
    max_c = min(30, ws.max_column)

    candidates = []
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                val_str = str(val).strip()
                if val_str.startswith('='):
                    continue
                val_norm = normalize_str(val_str)
                
                # Exclude header candidates that match total/summary words
                if any(k in val_norm for k in ["total", "somme", "sum", "totaux", "average", "moyenne"]):
                    continue
                    
                is_employee_cell = False
                if val_norm in EMPLOYEE_KEYWORDS:
                    is_employee_cell = True
                elif val_norm in {"nom", "name", "id", "employe", "employee"}:
                    is_employee_cell = True
                elif any(k in val_norm for k in ["employe", "employee", "collaborateur", "agent", "salarie", "staff", "personnel"]):
                    is_employee_cell = True
                    
                if is_employee_cell:
                    priority = 2 if any(k in val_norm for k in ["employe", "employee", "nom", "name", "collaborateur"]) else 1
                    candidates.append((r, c, priority))

    if candidates:
        candidates.sort(key=lambda x: (-x[2], x[0], x[1]))
        header_row, emp_col, _ = candidates[0]
    else:
        header_row = None
        emp_col = None
        for c in range(1, max_c + 1):
            text_count = 0
            num_count = 0
            first_text_row = None
            for r in range(1, min(15, ws.max_row) + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None:
                    if isinstance(val, (int, float)):
                        num_count += 1
                    elif isinstance(val, str) and val.strip():
                        if val.strip().startswith('='):
                            num_count += 1
                        else:
                            text_count += 1
                            if first_text_row is None:
                                first_text_row = r
            if text_count > 2 and text_count > num_count:
                header_row = first_text_row
                emp_col = c
                break

        if header_row is None:
            header_row = 1
            emp_col = 1

    # Detect prenom_col (first name column) if separate from emp_col (last name)
    prenom_col = None
    if header_row is not None:
        for c in range(1, max_c + 1):
            val = ws.cell(row=header_row, column=c).value
            if val is not None:
                val_norm = normalize_str(val)
                if "prenom" in val_norm:
                    prenom_col = c
                    break
        
        if prenom_col is not None and emp_col == prenom_col:
            found_other = False
            for c in range(1, max_c + 1):
                if c == prenom_col:
                    continue
                val = ws.cell(row=header_row, column=c).value
                if val is not None:
                    val_norm = normalize_str(val)
                    if any(k in val_norm for k in ["nom", "name", "employe", "employee", "collaborateur", "agent", "salarie"]):
                        emp_col = c
                        found_other = True
                        break
            if not found_other:
                prenom_col = None

    # Find employee rows and summary rows
    emp_rows = []
    total_rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        emp_val = ws.cell(row=r, column=emp_col).value
        if emp_val is None:
            continue
        
        val_str = str(emp_val).strip()
        if not val_str:
            continue
            
        if val_str.startswith('='):
            val_norm = val_str.lower()
            is_summary = any(k in val_norm for k in SUMMARY_KEYWORDS)
            if is_summary:
                total_rows.append(r)
            else:
                emp_rows.append(r)
        else:
            emp_str = normalize_str(emp_val)
            if emp_str in SUMMARY_KEYWORDS:
                total_rows.append(r)
            else:
                emp_rows.append(r)

    # Find columns to sum
    sum_cols = []
    existing_total_col = None

    for c in range(emp_col + 1, ws.max_column + 1):
        if prenom_col is not None and c == prenom_col:
            continue

        header_val = ws.cell(row=header_row, column=c).value
        header_norm = normalize_str(header_val)
        
        if header_norm in {"total", "somme", "sum", "totaux"}:
            existing_total_col = c
            continue
            
        if header_norm in EXCLUDED_HEADERS:
            continue
            
        # Verify if column is mostly numeric
        text_count = 0
        num_count = 0
        for r in emp_rows:
            val = ws.cell(row=r, column=c).value
            if val is not None:
                if isinstance(val, (int, float)):
                    num_count += 1
                elif isinstance(val, str):
                    val_str = val.strip()
                    if val_str.startswith('='):
                        num_count += 1
                    elif val_str:
                        text_count += 1
                        
        if text_count > 0 and text_count >= num_count:
            continue
            
        sum_cols.append(c)

    return header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col


def is_summary_sheet(ws, header_row, sum_cols):
    presence_col = find_column_by_header(ws, header_row, {'presence'})
    if presence_col:
        return False
        
    months = {"janvier", "fevrier", "mars", "avril", "mai", "juin", 
              "juillet", "aout", "septembre", "octobre", "novembre", "decembre"}
    
    match_count = 0
    for c in sum_cols:
        val = normalize_str(ws.cell(row=header_row, column=c).value)
        if any(m in val for m in months):
            match_count += 1
            
    return match_count > 0

def name_columns_for_key(ws, header_row, emp_col, prenom_col):
    nom_col = None
    first_name_col = None

    for c in range(1, ws.max_column + 1):
        header_norm = normalize_str(ws.cell(row=header_row, column=c).value)
        if header_norm in {'nom', 'nom.'}:
            nom_col = c
        elif header_norm in {'prenom', 'prenom.'} or (header_norm.startswith('pr') and 'nom' in header_norm):
            first_name_col = c

    return nom_col or emp_col, first_name_col or prenom_col


def is_name_candidate(value):
    if value is None:
        return False
    value_str = str(value).strip()
    if not value_str or value_str.startswith('='):
        return False
    if re.fullmatch(r'[\d\s:./-]+', value_str):
        return False
    value_norm = normalize_str(value_str)
    if value_norm in {'total', 'somme', 'sum', 'totaux'}:
        return False
    return any(ch.isalpha() for ch in value_str)


def name_signature(name):
    tokens = re.findall(r'[a-z0-9]+', normalize_str(name))
    tokens = [token for token in tokens if not token.isdigit()]
    return ' '.join(sorted(tokens))


def employee_name_candidates(ws, row, header_row, emp_col, prenom_col):
    candidates = []

    def add_candidate(value):
        if is_name_candidate(value):
            candidate = normalize_str(value)
            if candidate and candidate not in candidates:
                candidates.append(candidate)

    nom_col, first_name_col = name_columns_for_key(ws, header_row, emp_col, prenom_col)
    nom = ws.cell(row=row, column=nom_col).value if nom_col is not None else None
    prenom = ws.cell(row=row, column=first_name_col).value if first_name_col is not None else None

    if is_name_candidate(prenom) and is_name_candidate(nom):
        add_candidate(format_employee_name(prenom, nom))
        add_candidate(f"{nom} {prenom}")
    elif is_name_candidate(prenom):
        add_candidate(prenom)
    elif is_name_candidate(nom):
        add_candidate(nom)

    for c in range(1, ws.max_column + 1):
        header_norm = normalize_str(ws.cell(row=header_row, column=c).value)
        if header_norm in {'date', 'jour', 'time', 'temps', 'heure', 'presence', 'quantite', 'total', 'prix'}:
            continue
        add_candidate(ws.cell(row=row, column=c).value)

    return candidates


def get_employee_key(ws, row, header_row, emp_col, prenom_col):
    matricule_col = find_column_by_header(ws, header_row, {'matricule', 'matricule.'})
    emp_no_col = find_column_by_header(ws, header_row, {'emp no', 'emp no.'})
    
    code_val = None
    if matricule_col:
        val = ws.cell(row=row, column=matricule_col).value
        if val is not None and str(val).strip():
            code_val = str(val).strip().lower()
    elif emp_no_col:
        val = ws.cell(row=row, column=emp_no_col).value
        if val is not None and str(val).strip():
            code_val = str(val).strip().lower()

    candidates = employee_name_candidates(ws, row, header_row, emp_col, prenom_col)
    name_val = candidates[0] if candidates else ''
    
    return (code_val, name_val)


def quote_sheet_formula_name(sheet_name):
    escaped_name = sheet_name.replace("'", "''")
    return f"'{escaped_name}'"

def find_employee_cell(sheet_cells, emp_key):
    code, name = emp_key
    if code:
        for (c, n), coord in sheet_cells.items():
            if c == code:
                return coord
    if name:
        for (c, n), coord in sheet_cells.items():
            if n == name:
                return coord

        wanted_signature = name_signature(name)
        if wanted_signature:
            for (c, n), coord in sheet_cells.items():
                if name_signature(n) == wanted_signature:
                    return coord
    return None


def find_total_rows_anywhere(ws, header_row):
    total_headers = {"total", "somme", "sum", "totaux"}
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if normalize_str(ws.cell(row=r, column=c).value) in total_headers:
                rows.append(r)
                break
    return rows

def process_workbook_in_place(wb):
    sheet_employee_cells = {} # sheet_name -> { emp_key -> cell_coordinate }
    summary_sheets = [] # list of (ws, header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col)

    # First Pass: Process monthly sheets and collect employee cells
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col = detect_structure(ws)
        
        if not emp_rows or not sum_cols:
            continue

        # Check if summary sheet
        if is_summary_sheet(ws, header_row, sum_cols):
            summary_sheets.append((ws, header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col))
            continue
            
        sheet_employee_cells[sheet_name] = {}
        
        # If it's an attendance quantity sheet
        if process_attendance_quantity_sheet(ws, header_row, emp_rows, emp_col, prenom_col, sheet_employee_cells[sheet_name]):
            continue
            
        # Determine layout type
        name_counts = {}
        for r in emp_rows:
            val = ws.cell(row=r, column=emp_col).value
            p_val = ws.cell(row=r, column=prenom_col).value if prenom_col is not None else None
            key = (str(val).strip() if val is not None else "", str(p_val).strip() if p_val is not None else "")
            if key[0] or key[1]:
                name_counts[key] = name_counts.get(key, 0) + 1
                
        is_list_layout = False
        if len(emp_rows) > 0:
            max_occ = max(name_counts.values()) if name_counts else 0
            if max_occ > 1:
                is_list_layout = True
                
        if is_list_layout:
            # Collapse duplicates and sum values in sum_cols
            unique_emps = []
            emp_to_rows = {}
            for r in emp_rows:
                val = ws.cell(row=r, column=emp_col).value
                p_val = ws.cell(row=r, column=prenom_col).value if prenom_col is not None else None
                if val is not None or p_val is not None:
                    nom_str = str(val).strip() if val is not None else ""
                    prenom_str = str(p_val).strip() if p_val is not None else ""
                    key = (nom_str, prenom_str)
                    if key not in emp_to_rows:
                        unique_emps.append(key)
                        emp_to_rows[key] = []
                    emp_to_rows[key].append(r)
            
            rows_to_delete = []
            for key in unique_emps:
                rows = emp_to_rows[key]
                target_row = rows[0]
                
                # Sum daily columns
                for c in sum_cols:
                    col_sum = 0
                    has_numeric = False
                    for r in rows:
                        v = ws.cell(row=r, column=c).value
                        if v is not None:
                            if isinstance(v, (int, float)):
                                col_sum += v
                                has_numeric = True
                            elif isinstance(v, str):
                                val_str = v.strip()
                                if val_str.isdigit():
                                    col_sum += int(val_str)
                                    has_numeric = True
                    if has_numeric:
                        ws.cell(row=target_row, column=c, value=col_sum)
                
                # Collect duplicate rows to delete
                if len(rows) > 1:
                    rows_to_delete.extend(rows[1:])
            
            # Delete duplicates from bottom to top
            for r in sorted(rows_to_delete, reverse=True):
                ws.delete_rows(r)
                
            # Re-detect structure on the collapsed sheet
            header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col = detect_structure(ws)
            if not emp_rows or not sum_cols:
                continue

        # Find all existing total columns
        total_cols = []
        for c in range(max(emp_col, prenom_col or 0) + 1, ws.max_column + 1):
            header_val = ws.cell(row=header_row, column=c).value
            header_norm = normalize_str(header_val)
            if header_norm in {"total", "somme", "sum", "totaux"}:
                total_cols.append(c)
                
        # Decide total column
        if total_cols:
            total_col = total_cols[0]
        else:
            total_col = ws.max_column + 1
            hdr_cell = ws.cell(row=header_row, column=total_col, value="Total")
            copy_style(ws.cell(row=header_row, column=sum_cols[-1]), hdr_cell)
            
        total_col_letter = get_column_letter(total_col)
        start_col_letter = get_column_letter(sum_cols[0])
        end_col_letter = get_column_letter(sum_cols[-1])
        
        # Populate totals for employees
        for r in emp_rows:
            row_sum = 0
            for c in sum_cols:
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    if isinstance(v, (int, float)):
                        row_sum += v
                    elif isinstance(v, str):
                        val_str = v.strip()
                        if val_str.isdigit():
                            row_sum += int(val_str)
            cell = ws.cell(row=r, column=total_col, value=row_sum)
            copy_style(ws.cell(row=r, column=sum_cols[-1]), cell)
            
            # Save cell coordinate
            emp_key = get_employee_key(ws, r, header_row, emp_col, prenom_col)
            sheet_employee_cells[sheet_name][emp_key] = cell.coordinate
            
        # Populate totals for summary rows
        if emp_rows:
            for r in total_rows:
                col_sum = 0
                for er in emp_rows:
                    v = ws.cell(row=er, column=total_col).value
                    if isinstance(v, (int, float)):
                        col_sum += v
                cell = ws.cell(row=r, column=total_col, value=col_sum)
                copy_style(ws.cell(row=r, column=sum_cols[-1]), cell)

    # Second Pass: Process summary sheets and populate calculated values
    for ws, header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col in summary_sheets:
        # Find all existing total columns
        total_cols = []
        for c in range(max(emp_col, prenom_col or 0) + 1, ws.max_column + 1):
            header_val = ws.cell(row=header_row, column=c).value
            header_norm = normalize_str(header_val)
            if header_norm in {"total", "somme", "sum", "totaux"}:
                total_cols.append(c)
                
        if total_cols:
            total_col = total_cols[0]
        else:
            total_col = ws.max_column + 1
            hdr_cell = ws.cell(row=header_row, column=total_col, value="Total")
            copy_style(ws.cell(row=header_row, column=sum_cols[-1]), hdr_cell)
            
        # Fill month columns with already calculated values from the matching monthly sheets
        for c in sum_cols:
            month_header = normalize_str(ws.cell(row=header_row, column=c).value)
            
            # Find the matching sheet name
            matching_sheet_name = None
            for name in sheet_employee_cells.keys():
                if month_header in normalize_str(name):
                    matching_sheet_name = name
                    break
                    
            if matching_sheet_name:
                source_ws = wb[matching_sheet_name]
                for r in emp_rows:
                    emp_key = get_employee_key(ws, r, header_row, emp_col, prenom_col)
                    cell_coord = find_employee_cell(sheet_employee_cells[matching_sheet_name], emp_key)
                    if cell_coord:
                        source_value = numeric_value(source_ws[cell_coord].value)
                        ws.cell(row=r, column=c, value=safe_int_or_float(source_value))
                    else:
                        ws.cell(row=r, column=c, value=0)
            else:
                # If no matching sheet name found, fill with 0
                for r in emp_rows:
                    ws.cell(row=r, column=c, value=0)
                    
        # Fill row totals with calculated values
        price_col = find_column_by_header(ws, header_row, {'prix', 'price', 'valeur', 'montant'})
        total_rows = sorted(set(total_rows) | set(find_total_rows_anywhere(ws, header_row)))
        for r in emp_rows:
            row_total = sum(numeric_value(ws.cell(row=r, column=c).value) for c in sum_cols)
            cell = ws.cell(row=r, column=total_col, value=safe_int_or_float(row_total))
            copy_style(ws.cell(row=r, column=sum_cols[-1]), cell)

            if price_col and price_col != total_col:
                price_cell = ws.cell(row=r, column=price_col, value=safe_int_or_float(row_total * MILK_UNIT_PRICE))
                copy_style(cell, price_cell)
            

        for r in total_rows:
            for c in sum_cols:
                col_total = sum(numeric_value(ws.cell(row=er, column=c).value) for er in emp_rows)
                cell = ws.cell(row=r, column=c, value=safe_int_or_float(col_total))
                copy_style(ws.cell(row=r - 1, column=c), cell)

            summary_total = sum(numeric_value(ws.cell(row=r, column=c).value) for c in sum_cols)
            total_cell = ws.cell(row=r, column=total_col, value=safe_int_or_float(summary_total))
            copy_style(ws.cell(row=r - 1, column=total_col), total_cell)

            if price_col and price_col != total_col:
                price_total = sum(numeric_value(ws.cell(row=er, column=price_col).value) for er in emp_rows)
                price_cell = ws.cell(row=r, column=price_col, value=safe_int_or_float(price_total))
                copy_style(ws.cell(row=r - 1, column=price_col), price_cell)
            
    return wb


def lire_excel_dynamique(file_bytes):
    wb = load_workbook_from_bytes(file_bytes, data_only=True)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col = detect_structure(ws)
        if emp_rows and sum_cols:
            name_counts = {}
            for r in emp_rows:
                val = ws.cell(row=r, column=emp_col).value
                p_val = ws.cell(row=r, column=prenom_col).value if prenom_col is not None else None
                key = (str(val).strip() if val is not None else "", str(p_val).strip() if p_val is not None else "")
                if key[0] or key[1]:
                    name_counts[key] = name_counts.get(key, 0) + 1
                    
            is_list_layout = False
            if len(emp_rows) > 0:
                max_occ = max(name_counts.values()) if name_counts else 0
                if max_occ > 1:
                    is_list_layout = True
                    
            if not is_list_layout:
                dates = []
                for c in sum_cols:
                    hdr_val = ws.cell(row=header_row, column=c).value
                    dates.append(str(hdr_val) if hdr_val is not None else "")
                    
                employees = []
                data = {}
                for r in emp_rows:
                    nom = ws.cell(row=r, column=emp_col).value
                    prenom = ws.cell(row=r, column=prenom_col).value if prenom_col is not None else None
                    if nom is None and prenom is None:
                        continue
                    
                    nom_formatted = format_employee_name(prenom, nom)
                    if not nom_formatted:
                        continue
                        
                    if nom_formatted not in data:
                        employees.append(nom_formatted)
                        data[nom_formatted] = []
                    
                    vals = []
                    for c in sum_cols:
                        v = ws.cell(row=r, column=c).value
                        try:
                            vals.append(int(v) if v is not None else 0)
                        except:
                            vals.append(0)
                    data[nom_formatted] = vals
                return employees, dates, data
            else:
                dates = []
                for c in sum_cols:
                    hdr_val = ws.cell(row=header_row, column=c).value
                    dates.append(str(hdr_val) if hdr_val is not None else "")
                    
                employees_map = {}
                for r in emp_rows:
                    nom = ws.cell(row=r, column=emp_col).value
                    prenom = ws.cell(row=r, column=prenom_col).value if prenom_col is not None else None
                    if nom is None and prenom is None:
                        continue
                        
                    nom_formatted = format_employee_name(prenom, nom)
                    if not nom_formatted:
                        continue
                        
                    if nom_formatted not in employees_map:
                        employees_map[nom_formatted] = [0] * len(sum_cols)
                        
                    for idx, c in enumerate(sum_cols):
                        v = ws.cell(row=r, column=c).value
                        try:
                            employees_map[nom_formatted][idx] += int(v) if v is not None else 0
                        except:
                            pass
                            
                employees = list(employees_map.keys())
                return employees, dates, employees_map
                
    return [], [], {}

MONTH_NAMES = {
    1: "Janvier", 2: "FÃ©vrier", 3: "Mars", 4: "Avril", 5: "Mai", 6: "Juin",
    7: "Juillet", 8: "AoÃ»t", 9: "Septembre", 10: "Octobre", 11: "Novembre", 12: "DÃ©cembre"
}

def extraire_mois_annee(ws, dates):
    # 1. Try to find date in sheet title
    title = ws.title
    title_clean = normalize_str(title)
    
    # Check for French month names in sheet title
    for m_num, m_name in MONTH_NAMES.items():
        if normalize_str(m_name) in title_clean:
            year_match = re.search(r'\b(20\d{2})\b', title)
            year = year_match.group(1) if year_match else str(datetime.now().year)
            return f"{m_name} {year}"
            
    # Look for MM/YYYY or MM-YYYY in title
    match = re.search(r'\b(0?[1-9]|1[0-2])[-/](20\d{2})\b', title)
    if match:
        m_num = int(match.group(1))
        year = match.group(2)
        return f"{MONTH_NAMES[m_num]} {year}"

    # 2. Try to find from dates headers
    for d in dates:
        if isinstance(d, datetime):
            return f"{MONTH_NAMES[d.month]} {d.year}"
        if isinstance(d, str):
            # Try to match formats like DD/MM/YYYY or YYYY-MM-DD
            match_slash = re.search(r'\b\d{1,2}/(\d{1,2})/(20\d{2})\b', d)
            if match_slash:
                m_num = int(match_slash.group(1))
                year = match_slash.group(2)
                if m_num in MONTH_NAMES:
                    return f"{MONTH_NAMES[m_num]} {year}"
            
            match_dash = re.search(r'\b(20\d{2})-(\d{1,2})-\d{1,2}\b', d)
            if match_dash:
                year = match_dash.group(1)
                m_num = int(match_dash.group(2))
                if m_num in MONTH_NAMES:
                    return f"{MONTH_NAMES[m_num]} {year}"

    # 3. Try to look at the first few rows/cells of the sheet for any cell containing a month or date
    for r in range(1, min(10, ws.max_row + 1)):
        for c in range(1, min(10, ws.max_column + 1)):
            val = ws.cell(row=r, column=c).value
            if isinstance(val, datetime):
                return f"{MONTH_NAMES[val.month]} {val.year}"
            if isinstance(val, str):
                val_norm = normalize_str(val)
                for m_num, m_name in MONTH_NAMES.items():
                    if normalize_str(m_name) in val_norm:
                        year_match = re.search(r'\b(20\d{2})\b', val)
                        year = year_match.group(1) if year_match else str(datetime.now().year)
                        return f"{m_name} {year}"

    # 4. Fallback to current month and year
    now = datetime.now()
    return f"{MONTH_NAMES[now.month]} {now.year}"

MILK_UNIT_PRICE = 1350


def format_label_name(prenom_val, nom_val):
    nom = str(nom_val).strip() if nom_val is not None else ''
    prenom = str(prenom_val).strip() if prenom_val is not None else ''
    if nom and prenom:
        return f"{nom} {prenom}"
    if nom:
        return format_single_name(nom)
    if prenom:
        return format_single_name(prenom)
    return ''


def label_name_columns(ws, header_row, emp_col, prenom_col):
    nom_col = None
    first_name_col = None

    for c in range(1, ws.max_column + 1):
        header_norm = normalize_str(ws.cell(row=header_row, column=c).value)
        if header_norm in {'nom', 'nom.'}:
            nom_col = c
        elif header_norm in {'prenom', 'prenom.'} or (header_norm.startswith('pr') and 'nom' in header_norm):
            first_name_col = c

    return nom_col or emp_col, first_name_col or prenom_col


def label_name_for_row(ws, row, nom_col, prenom_col):
    nom = ws.cell(row=row, column=nom_col).value if nom_col is not None else None
    prenom = ws.cell(row=row, column=prenom_col).value if prenom_col is not None else None
    return format_label_name(prenom, nom)

def employee_label_key(ws, row, header_row, emp_col, prenom_col):
    matricule_col = find_column_by_header(ws, header_row, {'matricule', 'matricule.'})
    emp_no_col = find_column_by_header(ws, header_row, {'emp no', 'emp no.'})
    if matricule_col:
        value = ws.cell(row=row, column=matricule_col).value
        if value is not None and str(value).strip():
            return ('matricule', str(value).strip())
    if emp_no_col:
        value = ws.cell(row=row, column=emp_no_col).value
        if value is not None and str(value).strip():
            return ('emp_no', str(value).strip())

    nom_col, first_name_col = label_name_columns(ws, header_row, emp_col, prenom_col)
    return ('name', normalize_str(label_name_for_row(ws, row, nom_col, first_name_col)))


def add_label_total(labels_map, order, key, name, amount):
    if not name or amount == 0:
        return
    if key not in labels_map:
        order.append(key)
        labels_map[key] = {'nom': name, 'total': 0}
    labels_map[key]['total'] += amount


def generer_etiquettes_pdf(employees_data):
    page_width, page_height = A4
    cols = 2
    rows = 4
    receipts_per_page = cols * rows
    margin_x = 14
    margin_y = 18
    receipt_width = (page_width - margin_x * 2) / cols
    receipt_height = (page_height - margin_y * 2) / rows

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    border_color = HexColor('#bfbfbf')
    text_color = HexColor('#111111')

    def fit_bold_text(text, max_width, start_size=12, min_size=8):
        size = start_size
        while c.stringWidth(text, 'Helvetica-Bold', size) > max_width and size > min_size:
            size -= 0.5
        return size

    for index, emp in enumerate(employees_data):
        page_idx = index % receipts_per_page
        col_idx = page_idx % cols
        row_idx = page_idx // cols
        x = margin_x + col_idx * receipt_width
        y = page_height - margin_y - (row_idx + 1) * receipt_height

        c.setStrokeColor(border_color)
        c.setLineWidth(0.75)
        c.rect(x, y, receipt_width, receipt_height, stroke=True, fill=False)

        title = 'Reçu de partage de lait'
        c.setFillColor(text_color)
        c.setFont('Helvetica-Bold', 16)
        c.drawCentredString(x + receipt_width / 2, y + receipt_height - 24, title)

        left = x + 8
        line_y = y + receipt_height - 58
        name = emp['nom']
        total = safe_int_or_float(emp['total'])
        value = int(total * MILK_UNIT_PRICE)

        rows_text = [
            f"Code :{index + 1}",
            f"Nom &Prénom :{name}",
            f"Quantité : {total}",
            f"Valeur financière : {value}",
        ]

        for text_line in rows_text:
            font_size = fit_bold_text(text_line, receipt_width - 18, start_size=12, min_size=8)
            if text_line.startswith('Valeur'):
                font_size = fit_bold_text(text_line, receipt_width - 18, start_size=13, min_size=9)
            c.setFont('Helvetica-Bold', font_size)
            c.drawString(left, line_y, text_line)
            line_y -= 32

        if page_idx == receipts_per_page - 1 and index < len(employees_data) - 1:
            c.showPage()

    c.save()
    buffer.seek(0)
    return buffer.getvalue()

@app.route("/api/labels", methods=["POST"])
def labels():
    """Génère et retourne le PDF des reçus cumulés par employé."""
    if "file" not in request.files:
        return jsonify({"error": "Aucun fichier reçu"}), 400
    f = request.files["file"]
    try:
        file_bytes = f.read()
        wb = load_workbook_from_bytes(file_bytes, data_only=True)

        labels_map = {}
        order = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_row, emp_col, prenom_col, emp_rows, total_rows, sum_cols, existing_total_col = detect_structure(ws)
            if not emp_rows:
                continue

            nom_col, first_name_col = label_name_columns(ws, header_row, emp_col, prenom_col)

            presence_col = find_column_by_header(ws, header_row, {'presence'})
            if presence_col:
                for r in emp_rows:
                    name = label_name_for_row(ws, r, nom_col, first_name_col)
                    key = employee_label_key(ws, r, header_row, emp_col, prenom_col)
                    add_label_total(labels_map, order, key, name, numeric_value(ws.cell(row=r, column=presence_col).value))
                continue

            if not sum_cols:
                continue

            name_counts = {}
            for r in emp_rows:
                val = ws.cell(row=r, column=nom_col).value
                p_val = ws.cell(row=r, column=first_name_col).value if first_name_col is not None else None
                key = (str(val).strip() if val is not None else '', str(p_val).strip() if p_val is not None else '')
                if key[0] or key[1]:
                    name_counts[key] = name_counts.get(key, 0) + 1

            is_list_layout = bool(name_counts and max(name_counts.values()) > 1)

            if is_list_layout:
                for r in emp_rows:
                    name = label_name_for_row(ws, r, nom_col, first_name_col)
                    key = employee_label_key(ws, r, header_row, emp_col, prenom_col)
                    row_total = sum(numeric_value(ws.cell(row=r, column=c).value) for c in sum_cols)
                    add_label_total(labels_map, order, key, name, row_total)
            else:
                for r in emp_rows:
                    name = label_name_for_row(ws, r, nom_col, first_name_col)
                    key = employee_label_key(ws, r, header_row, emp_col, prenom_col)
                    row_total = sum(numeric_value(ws.cell(row=r, column=c).value) for c in sum_cols)
                    add_label_total(labels_map, order, key, name, row_total)

        labels_data = [labels_map[key] for key in order if labels_map[key]['total'] > 0]

        if not labels_data:
            return jsonify({"error": "Aucune donnée d'employé trouvée dans le fichier Excel."}), 400

        pdf_bytes = generer_etiquettes_pdf(labels_data)

        response = make_response(pdf_bytes)
        response.headers["Content-Type"] = "application/pdf"
        response.headers["Content-Disposition"] = f"attachment; filename=recus_lait_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
        return response
    except Exception as e:
        details = ""
        if 'file_bytes' in locals():
            details = f" (Taille: {len(file_bytes)} octets, Début: {file_bytes[:10]})"
        error_msg = f"{str(e)}{details}"
        import sys
        print(f"ERROR labels: {error_msg}", file=sys.stderr)
        try:
            log_path = os.path.join(os.path.dirname(__file__), 'error.log')
            with open(log_path, 'a', encoding='utf-8') as logf:
                logf.write(f"[{datetime.now().isoformat()}] Exception in labels: {error_msg}\n")
                logf.write(traceback.format_exc())
                logf.write('\n')
        except Exception:
            pass
        return jsonify({"error": error_msg}), 500

@app.route("/api/preview", methods=["POST"])
def preview():
    """Upload un fichier, retourne les donnÃ©es JSON pour le dashboard."""
    if "file" not in request.files:
        return jsonify({"error": "Aucun fichier reÃ§u"}), 400
    f = request.files["file"]
    try:
        file_bytes = f.read()
        employees, dates, data = lire_excel_dynamique(file_bytes)

        # Calculs statistiques
        stats_employees = []
        total_global = 0
        for emp in employees:
            vals = data[emp]
            total = sum(vals)
            total_global += total
            stats_employees.append({
                "nom": emp,
                "total": total,
                "sans": len(dates) - total,
                "taux": round(total / len(dates) * 100, 1) if dates else 0,
                "valeurs": vals
            })

        stats_employees.sort(key=lambda x: x["total"], reverse=True)

        return jsonify({
            "employees": stats_employees,
            "dates": dates,
            "nb_dates": len(dates),
            "nb_employees": len(employees),
            "total_global": total_global,
            "taux_global": round(total_global / (len(employees) * len(dates)) * 100, 1) if employees and dates else 0
        })
    except Exception as e:
        try:
            log_path = os.path.join(os.path.dirname(__file__), 'error.log')
            with open(log_path, 'a', encoding='utf-8') as logf:
                logf.write(f"[{datetime.now().isoformat()}] Exception in download:\n")
                logf.write(traceback.format_exc())
                logf.write('\n')
        except Exception:
            pass
        return jsonify({"error": str(e)}), 500

@app.route("/api/download", methods=["POST"])
def download():
    """GÃ©nÃ¨re et retourne le fichier Excel ou CSV rapport."""
    if "file" not in request.files:
        return jsonify({"error": "Aucun fichier reÃ§u"}), 400
    f = request.files["file"]
    try:
        file_bytes = f.read()
        is_csv = is_csv_bytes(file_bytes)
        
        # Load workbook
        wb = load_workbook_from_bytes(file_bytes)
        processed_wb = process_workbook_in_place(wb)
        
        if is_csv:
            # Export to CSV
            output = io.StringIO()
            writer = csv.writer(output, delimiter=';') # Use semicolon as standard for European Excel
            ws = processed_wb.active
            for row in ws.iter_rows(values_only=True):
                writer.writerow([val if val is not None else "" for val in row])
            
            csv_bytes = output.getvalue().encode('utf-8-sig') # UTF-8 with BOM for Excel compatibility
            response = make_response(csv_bytes)
            response.headers["Content-Type"] = "text/csv"
            response.headers["Content-Disposition"] = f"attachment; filename=rapport_total_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
            return response
        else:
            # Export to Excel
            buf = io.BytesIO()
            processed_wb.save(buf)
            excel_bytes = buf.getvalue()
            response = make_response(excel_bytes)
            response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            response.headers["Content-Disposition"] = f"attachment; filename=rapport_total_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
            return response
    except Exception as e:
        details = ""
        if 'file_bytes' in locals():
            details = f" (Taille: {len(file_bytes)} octets, Début: {file_bytes[:10]})"
        error_msg = f"{str(e)}{details}"
        import sys
        print(f"ERROR download: {error_msg}", file=sys.stderr)
        try:
            log_path = os.path.join(os.path.dirname(__file__), 'error.log')
            with open(log_path, 'a', encoding='utf-8') as logf:
                logf.write(f"[{datetime.now().isoformat()}] Exception in download: {error_msg}\n")
                logf.write(traceback.format_exc())
                logf.write('\n')
        except Exception:
            pass
        return jsonify({"error": error_msg}), 500


@app.route("/api/health")
def health():
    return jsonify({"status": "ok"})

if __name__ == "__main__":
    app.run(debug=True, port=5000)
